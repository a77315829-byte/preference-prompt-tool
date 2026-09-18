"""수행계획서용 pptx 렌더러.

HTML 판(plan.html)의 컴포넌트를 그대로 도형으로 옮긴다. 자동 흐름이 없으므로
모든 블록이 자기 높이를 스스로 계산하고, 쌓은 높이가 본문 영역을 넘으면
build 단계에서 경고를 낸다 - 39장을 눈으로 확인하지 않고도 넘침을 잡기 위한 장치다.
"""

from __future__ import annotations

import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ── 토큰 (HTML 판의 :root 와 동일) ──────────────────────────────────────
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x18, 0x1D)
INK2 = RGBColor(0x4A, 0x54, 0x60)
INK3 = RGBColor(0x7D, 0x87, 0x94)
ACCENT = RGBColor(0xB0, 0x12, 0x1A)
ACCENT_SOFT = RGBColor(0xF7, 0xE9, 0xEA)
RULE = RGBColor(0xD8, 0xDD, 0xE4)
FILL = RGBColor(0xF3, 0xF5, 0xF8)
OK = RGBColor(0x1E, 0x5E, 0x3C)
OK_SOFT = RGBColor(0xEA, 0xF3, 0xEE)
WARN = RGBColor(0x8A, 0x5A, 0x00)
WARN_SOFT = RGBColor(0xFB, 0xF2, 0xE2)
COVER_SUB = RGBColor(0xB9, 0xC1, 0xCC)
COVER_META = RGBColor(0x8E, 0x97, 0xA3)
COVER_KICK = RGBColor(0xFF, 0x6B, 0x74)
COVER_RULE = RGBColor(0x33, 0x3A, 0x44)

KR = "맑은 고딕"
MONO = "Consolas"

SLIDE_W, SLIDE_H = 13.333, 7.5
PAD_X = 0.62
BODY_TOP_DEFAULT = 1.95
BODY_BOTTOM = 7.02
CONTENT_W = SLIDE_W - 2 * PAD_X

FS_MEGA, FS_H1, FS_H2 = 40, 30, 20
FS_LEAD, FS_BODY, FS_SM, FS_XS = 14, 11, 9.5, 8
FS_NUM, FS_CAP = 24, 7.5

GAP = 0.13          # 블록 사이 간격
HAIR = 0.009        # 선 두께


def _lines(text: str, width_in: float, pt: float) -> int:
    """폭 안에 들어가는 줄 수 추정. 한글은 글자폭이 pt 에 가깝고 라틴은 절반이라
    혼용 텍스트 기준 계수 0.60 을 실측으로 쓴다."""
    cpl = max(6, int(width_in * 72 / (pt * 0.60)))
    total = 0
    for para in text.split("\n"):
        total += max(1, math.ceil(len(para) / cpl))
    return total


def text_h(text: str, width_in: float, pt: float, lh: float = 1.45) -> float:
    return _lines(text, width_in, pt) * pt * lh / 72


def plain(runs) -> str:
    """서식 런 목록을 길이 추정용 평문으로 만든다."""
    if isinstance(runs, str):
        return runs
    if runs and all(isinstance(r, str) for r in runs):
        return "".join(runs)
    return "".join(r[0] for r in runs)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []

    # ── 원시 도형 ──────────────────────────────────────────────────
    def _slide(self):
        return self.prs.slides.add_slide(self.blank)

    def rect(self, s, x, y, w, h, fill=None, line=None, lw=0.75):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(lw)
        sh.text_frame.text = ""
        return sh

    def text(self, s, x, y, w, h, runs, size=FS_BODY, color=INK2, font=KR,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             lh=1.45, space_after=0):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if isinstance(runs, str):
            runs = [(runs, {})]
        elif runs and all(isinstance(r, str) for r in runs):
            runs = [("".join(runs), {})]
        paras = [[]]
        for txt, fmt in runs:
            pieces = txt.split("\n")
            for i, piece in enumerate(pieces):
                if i:
                    paras.append([])
                if piece:
                    paras[-1].append((piece, fmt))
        for pi, para_runs in enumerate(paras):
            p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = lh
            p.space_after = Pt(space_after)
            if not para_runs:
                para_runs = [(" ", {})]
            for txt, fmt in para_runs:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = fmt.get("font", font)
                f.size = Pt(fmt.get("size", size))
                f.bold = fmt.get("bold", bold)
                f.color.rgb = fmt.get("color", color)
        return box

    # ── 장 구성 ────────────────────────────────────────────────────
    def page(self, chapter, title, lead=None, folio=None, body_top=None):
        s = self._slide()
        self.rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=PAPER)
        self.rect(s, 0, SLIDE_H - 0.04, 2.7, 0.04, fill=ACCENT)
        self.text(s, PAD_X, 0.21, 6.0, 0.22, "수행계획서", size=FS_XS,
                  color=INK3, font=MONO)
        self.text(s, SLIDE_W - PAD_X - 6.0, 0.21, 6.0, 0.22, chapter,
                  size=FS_XS, color=ACCENT, font=MONO, bold=True,
                  align=PP_ALIGN.RIGHT)
        self.rect(s, PAD_X, 0.58, 0.045, 0.40, fill=ACCENT)
        self.text(s, PAD_X + 0.18, 0.58, CONTENT_W - 0.18, 0.42, title,
                  size=FS_H2, color=INK, bold=True, lh=1.15)
        y = 1.14
        if lead:
            lh_ = text_h(plain(lead), 10.6, FS_LEAD, 1.38)
            self.text(s, PAD_X, y, 10.6, lh_ + 0.1, lead, size=FS_LEAD,
                      color=INK, bold=True, lh=1.38)
            y += lh_ + 0.30
        if folio:
            self.text(s, SLIDE_W - PAD_X - 1.2, 7.13, 1.2, 0.2, folio,
                      size=FS_XS, color=INK3, font=MONO, align=PP_ALIGN.RIGHT)
        return s, (body_top if body_top is not None else max(y, 1.80))

    def cover(self, kicker, h1, sub, meta):
        s = self._slide()
        self.rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=INK)
        self.rect(s, 0, SLIDE_H - 0.055, SLIDE_W, 0.055, fill=ACCENT)
        self.text(s, PAD_X + 0.3, 1.52, 10.0, 0.25, kicker, size=FS_SM,
                  color=COVER_KICK, font=MONO, bold=True)
        self.text(s, PAD_X + 0.3, 1.95, 11.0, 1.9, h1, size=FS_MEGA,
                  color=PAPER, bold=True, lh=1.08)
        self.text(s, PAD_X + 0.3, 3.95, 10.4, 1.0, sub, size=FS_LEAD,
                  color=COVER_SUB, lh=1.45)
        self.rect(s, PAD_X + 0.3, 5.60, 11.6, HAIR, fill=COVER_RULE)
        cw = 11.6 / len(meta)
        for i, (k, v) in enumerate(meta):
            x = PAD_X + 0.3 + i * cw
            self.text(s, x, 5.80, cw - 0.2, 0.22, k, size=FS_SM,
                      color=COVER_META)
            self.text(s, x, 6.08, cw - 0.2, 0.55, v, size=FS_SM,
                      color=PAPER, bold=True, lh=1.35)
        return s

    def chapter(self, chapter, no, title, items, folio):
        s = self._slide()
        self.rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=FILL)
        self.rect(s, 0, SLIDE_H - 0.04, 2.7, 0.04, fill=ACCENT)
        self.text(s, PAD_X, 0.21, 6.0, 0.22, "수행계획서", size=FS_XS,
                  color=INK3, font=MONO)
        self.text(s, SLIDE_W - PAD_X - 6.0, 0.21, 6.0, 0.22, chapter,
                  size=FS_XS, color=ACCENT, font=MONO, bold=True,
                  align=PP_ALIGN.RIGHT)
        self.text(s, PAD_X + 0.2, 2.05, 4.0, 1.5, no, size=72, color=ACCENT,
                  font=MONO, bold=True, lh=0.95)
        self.text(s, PAD_X + 0.2, 3.45, 9.0, 0.6, title, size=FS_H1,
                  color=INK, bold=True, lh=1.15)
        runs = []
        for i, it in enumerate(items):
            if i:
                runs.append(("     ", {}))
            runs.append(("— ", {"color": ACCENT, "bold": True}))
            runs.append((it, {"color": INK2}))
        self.text(s, PAD_X + 0.2, 4.35, 11.4, 0.9, runs, size=FS_BODY,
                  lh=1.7)
        self.text(s, SLIDE_W - PAD_X - 1.2, 7.13, 1.2, 0.2, folio,
                  size=FS_XS, color=INK3, font=MONO, align=PP_ALIGN.RIGHT)
        return s


# ── 블록 ────────────────────────────────────────────────────────────
NOTE_STYLES = {
    "note": (FILL, INK3),
    "red": (ACCENT_SOFT, ACCENT),
    "ok": (OK_SOFT, OK),
    "warn": (WARN_SOFT, WARN),
    "layer": (PAPER, None),
}


class Note:
    """제목 + 본문(또는 글머리 목록) 박스."""

    def __init__(self, kind="note", head=None, text=None, bullets=None,
                 tail=None, size=FS_SM):
        self.kind, self.head = kind, head
        self.text, self.bullets, self.tail = text, bullets, tail
        self.size = size

    def height(self, w):
        inner = w - 0.34
        h = 0.20
        if self.head:
            h += text_h(plain(self.head), inner, self.size, 1.3) + 0.08
        if self.text:
            h += text_h(plain(self.text), inner, self.size, 1.5)
        if self.bullets:
            for b in self.bullets:
                h += text_h(plain(b), inner - 0.16, self.size, 1.45) + 0.055
        if self.tail:
            h += 0.07 + text_h(plain(self.tail), inner, self.size, 1.5)
        return h

    def draw(self, d, s, x, y, w):
        h = self.height(w)
        fill, border = NOTE_STYLES[self.kind]
        if self.kind == "layer":
            d.rect(s, x, y, w, h, fill=PAPER, line=RULE)
        else:
            d.rect(s, x, y, w, h, fill=fill)
            d.rect(s, x, y, 0.035, h, fill=border)
        inner = w - 0.34
        ty = y + 0.10
        if self.head:
            hh = text_h(plain(self.head), inner, self.size, 1.3)
            d.text(s, x + 0.22, ty, inner, hh + 0.05, self.head,
                   size=self.size, color=INK, bold=True, lh=1.3)
            ty += hh + 0.08
        if self.text:
            th = text_h(plain(self.text), inner, self.size, 1.5)
            d.text(s, x + 0.22, ty, inner, th + 0.05, self.text,
                   size=self.size, color=INK2, lh=1.5)
            ty += th
        if self.bullets:
            for b in self.bullets:
                bh = text_h(plain(b), inner - 0.16, self.size, 1.45)
                d.rect(s, x + 0.24, ty + self.size * 0.010, 0.055, 0.055,
                       fill=ACCENT)
                d.text(s, x + 0.40, ty, inner - 0.16, bh + 0.05, b,
                       size=self.size, color=INK2, lh=1.45)
                ty += bh + 0.055
        if self.tail:
            ty += 0.07
            th = text_h(plain(self.tail), inner, self.size, 1.5)
            d.text(s, x + 0.22, ty, inner, th + 0.05, self.tail,
                   size=self.size, color=INK2, lh=1.5)
        return h


class Table:
    def __init__(self, head, rows, weights, caption=None, aligns=None,
                 size=FS_SM):
        self.head, self.rows, self.weights = head, rows, weights
        self.caption, self.size = caption, size
        self.aligns = aligns or ["l"] * len(head)

    def _cols(self, w):
        total = sum(self.weights)
        return [w * v / total for v in self.weights]

    def _row_h(self, cells, cols, size, bold=False):
        hi = 0.0
        for cell, cw in zip(cells, cols):
            txt = plain(cell) if not isinstance(cell, tuple) else plain(cell[0])
            hi = max(hi, text_h(txt, cw - 0.18, size, 1.35))
        return hi + 0.135

    def height(self, w):
        cols = self._cols(w)
        h = 0.0
        if self.caption:
            h += text_h(self.caption, w, FS_CAP, 1.3) + 0.09
        h += self._row_h(self.head, cols, FS_XS) + HAIR
        for r in self.rows:
            h += self._row_h(r, cols, self.size) + HAIR
        return h

    def draw(self, d, s, x, y, w):
        cols = self._cols(w)
        y0 = y
        if self.caption:
            ch = text_h(self.caption, w, FS_CAP, 1.3)
            d.text(s, x, y, w, ch + 0.05, self.caption.upper(), size=FS_CAP,
                   color=INK3, font=MONO, lh=1.3)
            y += ch + 0.09
        hh = self._row_h(self.head, cols, FS_XS)
        d.rect(s, x, y, w, hh, fill=FILL)
        cx = x
        for cell, cw, al in zip(self.head, cols, self.aligns):
            d.text(s, cx + 0.09, y + 0.06, cw - 0.18, hh, cell, size=FS_XS,
                   color=INK, bold=True, lh=1.3,
                   align=PP_ALIGN.RIGHT if al == "r" else PP_ALIGN.LEFT)
            cx += cw
        y += hh
        d.rect(s, x, y, w, HAIR, fill=INK3)
        y += HAIR
        for ri, r in enumerate(self.rows):
            rh = self._row_h(r, cols, self.size)
            cx = x
            for cell, cw, al in zip(r, cols, self.aligns):
                runs, color, bold, font = cell, INK2, False, KR
                if isinstance(cell, tuple):
                    runs, style = cell[0], cell[1]
                    color = style.get("color", INK2)
                    bold = style.get("bold", False)
                    font = style.get("font", KR)
                elif al == "r":
                    font = MONO
                    color = INK
                d.text(s, cx + 0.09, y + 0.055, cw - 0.18, rh, runs,
                       size=self.size, color=color, bold=bold, font=font,
                       lh=1.35,
                       align=PP_ALIGN.RIGHT if al == "r" else PP_ALIGN.LEFT)
                cx += cw
            y += rh
            last = ri == len(self.rows) - 1
            d.rect(s, x, y, w, HAIR, fill=INK3 if last else RULE)
            y += HAIR
        return y - y0


class Stats:
    """상단 액센트 선 + 큰 수치 + 설명."""

    def __init__(self, items, cols=None, size=FS_NUM):
        self.items, self.cols, self.size = items, cols or len(items), size

    def _cell_w(self, w):
        return (w - (self.cols - 1) * 0.22) / self.cols

    def _row_h(self, w):
        cw = self._cell_w(w)
        lab = max(text_h(l, cw, FS_XS, 1.35) for _, l in self.items)
        return 0.06 + self.size * 1.05 / 72 + 0.09 + lab

    def height(self, w):
        rows = math.ceil(len(self.items) / self.cols)
        return rows * self._row_h(w) + (rows - 1) * 0.22

    def draw(self, d, s, x, y, w):
        """cols 개씩 줄바꿈해 격자로 놓는다. 1열로 두면 세로로 쌓인다."""
        cw = self._cell_w(w)
        rh = self._row_h(w)
        for i, (val, lab) in enumerate(self.items):
            cx = x + (i % self.cols) * (cw + 0.22)
            cy = y + (i // self.cols) * (rh + 0.22)
            d.rect(s, cx, cy, cw, 0.032, fill=ACCENT)
            d.text(s, cx, cy + 0.12, cw, self.size * 1.1 / 72 + 0.1, val,
                   size=self.size, color=INK, font=MONO, bold=True, lh=1.0)
            d.text(s, cx, cy + 0.12 + self.size * 1.05 / 72 + 0.09, cw,
                   rh, lab, size=FS_XS, color=INK3, lh=1.35)
        return self.height(w)


class Flow:
    def __init__(self, items, core=None):
        self.items, self.core = items, core

    def height(self, w):
        n = len(self.items)
        bw = (w - (n - 1) * 0.30) / n
        body = max(text_h(t, bw - 0.24, FS_XS, 1.4) for _, t in self.items)
        return 0.14 + FS_SM * 1.3 / 72 + 0.06 + body + 0.14

    def draw(self, d, s, x, y, w):
        n = len(self.items)
        bw = (w - (n - 1) * 0.30) / n
        h = self.height(w)
        for i, (title, body) in enumerate(self.items):
            cx = x + i * (bw + 0.30)
            is_core = i == self.core
            d.rect(s, cx, y, bw, h, fill=ACCENT_SOFT if is_core else PAPER,
                   line=ACCENT if is_core else RULE,
                   lw=1.5 if is_core else 0.75)
            d.text(s, cx + 0.12, y + 0.12, bw - 0.24, FS_SM * 1.4 / 72 + 0.05,
                   title, size=FS_SM, color=INK, bold=True, lh=1.3)
            d.text(s, cx + 0.12, y + 0.12 + FS_SM * 1.3 / 72 + 0.06,
                   bw - 0.24, h, body, size=FS_XS, color=INK2, lh=1.4)
            if i < n - 1:
                d.text(s, cx + bw + 0.03, y + h / 2 - 0.12, 0.24, 0.24, "→",
                       size=FS_SM, color=INK3, font=MONO,
                       align=PP_ALIGN.CENTER)
        return h


class Bars:
    """(라벨, 표시값, 0~1 비율, 색) 목록."""

    def __init__(self, items, note=None, label_w=1.5, val_w=0.62):
        self.items, self.note = items, note
        self.label_w, self.val_w = label_w, val_w

    def height(self, w):
        h = len(self.items) * (0.215 + 0.075)
        if self.note:
            h += 0.06 + text_h(self.note, w, FS_XS, 1.35)
        return h

    def draw(self, d, s, x, y, w):
        track_x = x + self.label_w + 0.12
        track_w = w - self.label_w - self.val_w - 0.28
        y0 = y
        for lab, val, frac, color in self.items:
            d.text(s, x, y + 0.03, self.label_w, 0.2, lab, size=FS_XS,
                   color=INK2, align=PP_ALIGN.RIGHT)
            d.rect(s, track_x, y, track_w, 0.215, fill=FILL)
            d.rect(s, track_x, y, max(track_w * frac, 0.02), 0.215,
                   fill=color)
            d.text(s, track_x + track_w + 0.10, y + 0.03, self.val_w, 0.2,
                   val, size=FS_XS, color=INK, font=MONO,
                   align=PP_ALIGN.RIGHT)
            y += 0.215 + 0.075
        if self.note:
            y += 0.06
            d.text(s, x, y, w, text_h(self.note, w, FS_XS, 1.35) + 0.05,
                   self.note, size=FS_XS, color=INK3, lh=1.35)
        return y - y0


class Gantt:
    def __init__(self, weeks, rows, legend, label_w=3.1):
        self.weeks, self.rows, self.legend = weeks, rows, legend
        self.label_w = label_w

    def height(self, w):
        return 0.26 + len(self.rows) * 0.295 + 0.08 + 0.24

    def draw(self, d, s, x, y, w):
        n = len(self.weeks)
        grid_x = x + self.label_w
        cw = (w - self.label_w) / n
        for i, wk in enumerate(self.weeks):
            d.text(s, grid_x + i * cw, y, cw, 0.2, wk, size=FS_XS,
                   color=INK3, align=PP_ALIGN.CENTER)
        y += 0.26
        for label, spans, color, dim in self.rows:
            d.text(s, x, y + 0.045, self.label_w - 0.14, 0.24, label,
                   size=FS_XS, color=INK3 if dim else INK, lh=1.25)
            for i in range(n):
                on = any(a <= i + 1 <= b for a, b in spans)
                d.rect(s, grid_x + i * cw + 0.018, y + 0.055,
                       cw - 0.036, 0.145, fill=color if on else FILL)
            y += 0.295
            d.rect(s, x, y - 0.045, w, HAIR, fill=RULE)
        y += 0.08
        runs = []
        for i, (swatch, txt) in enumerate(self.legend):
            if i:
                runs.append(("      ", {}))
            if swatch is not None:
                runs.append(("■ ", {"color": swatch}))
            runs.append((txt, {"color": INK3}))
        d.text(s, x, y, w, 0.24, runs, size=FS_XS, lh=1.3)
        return self.height(w)


class Cols:
    """열마다 블록을 세로로 쌓는다."""

    def __init__(self, columns, weights=None, gap=0.26):
        self.columns, self.gap = columns, gap
        self.weights = weights or [1] * len(columns)

    def widths(self, w):
        avail = w - self.gap * (len(self.columns) - 1)
        total = sum(self.weights)
        return [avail * v / total for v in self.weights]

    def height(self, w):
        ws = self.widths(w)
        best = 0.0
        for col, cwid in zip(self.columns, ws):
            h = sum(b.height(cwid) for b in col) + GAP * max(0, len(col) - 1)
            best = max(best, h)
        return best

    def draw(self, d, s, x, y, w):
        ws = self.widths(w)
        cx = x
        for col, cwid in zip(self.columns, ws):
            cy = y
            for b in col:
                cy += b.draw(d, s, cx, cy, cwid) + GAP
            cx += cwid + self.gap
        return self.height(w)


def stack(d, s, blocks, top, x=PAD_X, w=CONTENT_W, label=""):
    """블록을 세로로 쌓고 넘치면 경고를 남긴다."""
    y = top
    for b in blocks:
        y += b.draw(d, s, x, y, w) + GAP
    over = y - GAP - BODY_BOTTOM
    if over > 0.01:
        d.warnings.append(f"{label}: 본문이 {over:.2f}in 넘침")
    return y
